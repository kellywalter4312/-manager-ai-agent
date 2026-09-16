"""Phase 3 completion — Files, Documents & Control Center.

CHECK found General Knowledge, Live Information, Research, PDF/text
support and the Control Center's own terminal REPL already real
(§52.22). This file proves what was added to close the remaining gap:

* Spreadsheet (.xlsx/.csv) and presentation (.pptx) support in
  ``DocumentKnowledgeProvider`` - real row/slide extraction, never OCR,
  alongside the existing, unmodified PDF/text support.
* Real, honestly-bounded image analysis - ``AnthropicProvider.
  analyze_image`` (a genuine, additive Claude vision call, never part
  of the core ``LLMProvider`` contract), reached through a new
  ``ManagerBrain.analyze_image`` question (bound the identical way
  ``converse`` already is) and the ``image_analysis`` capability -
  deliberately never a ``Tool``, since a ``Tool`` must never hold
  reasoning machinery (an early version of this did, and was caught
  and corrected by this project's own existing
  ``test_the_toolset_consults_no_brain`` architectural test before
  being shipped). Every provider without real vision support is
  refused honestly, at the Brain layer.
* The Control Center's own new ``capabilities``/``provider`` commands -
  the real capability matrix and the real, existing Ecosystem/
  CredentialStore management surface, reached from the same terminal a
  person already has open, never a second UI.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pytest
from pptx import Presentation

from src.agents.tools.toolset import ToolSet
from src.config.settings import Settings
from src.control_center.observer import ControlCenter
from src.core.contracts import Task
from src.manager.capabilities import image_analysis_intent, select_capability
from src.providers.anthropic_provider import AnthropicProvider
from src.providers.base import Completion, LLMProvider
from src.providers.document_knowledge_provider import DocumentKnowledgeProvider
from src.manager.llm_brain import LLMManagerBrain
from src.services.ecosystem.contracts import EcosystemRegistry, GenericCategoryAdapter
from src.services.ecosystem.registry_factory import build_ecosystem_registry
from scripts.control_center import build, render_capabilities, run_command, run_provider_command
from scripts.exchange import exchange


def real_system(**overrides):
    return build(Settings.from_env(env={}, memory="ephemeral", **overrides))


def _with_brain(runtime, manager, text: str):
    class _Fixed(LLMProvider):
        @property
        def name(self):
            return "fixed"

        def complete(self, prompt, *, system=None):
            return Completion(text=text, provider="fixed")

    manager.brain = LLMManagerBrain(_Fixed(), personality=runtime.identity.personality)
    manager.own_skills._converse = manager.brain.converse


def _with_vision_brain(runtime, manager, provider: LLMProvider):
    """Bind a real LLMManagerBrain wrapping ``provider`` so
    ``request.analyze_image`` reaches it - the identical wiring
    ``composition.py`` does in production
    (``analyze_image=runtime.manager_brain.analyze_image``), reproduced
    directly for a test double."""
    manager.brain = LLMManagerBrain(provider, personality=runtime.identity.personality)
    manager.own_skills._converse = manager.brain.converse
    manager.own_skills._analyze_image = manager.brain.analyze_image


class _FakeVisionProvider(LLMProvider):
    """A real LLMProvider double that also exposes analyze_image -
    proving the mechanism, never claiming a live vendor call."""

    def __init__(self, text: str = "I see a red square.") -> None:
        self._text = text
        self.calls: list[tuple[bytes, str, str]] = []

    @property
    def name(self) -> str:
        return "fake-vision"

    def complete(self, prompt, *, system=None):
        return Completion(text="n/a", provider=self.name)

    def analyze_image(self, image_bytes: bytes, prompt: str, *, media_type: str = "image/png"):
        self.calls.append((image_bytes, prompt, media_type))
        return Completion(text=self._text, provider=self.name)


def _real_png(color=(255, 0, 0)) -> bytes:
    import io

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), color=color).save(buffer, format="PNG")
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 1. Spreadsheet support - real rows, real search, real analysis
# ---------------------------------------------------------------------------


class TestSpreadsheetSupport:
    def test_xlsx_rows_are_extracted_with_header_labelled_cells(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Name", "Revenue"])
        ws.append(["Alice", 4200])
        ws.append(["Bob", 9100])
        wb.save(tmp_path / "sales.xlsx")

        provider = DocumentKnowledgeProvider(tmp_path)
        result = provider.lookup("Bob revenue")
        assert result.excerpts
        assert any("Bob" in e.content and "9100" in e.content for e in result.excerpts)
        assert all(e.locator.startswith("row") for e in result.excerpts)

    def test_csv_needs_no_extra_dependency(self, tmp_path):
        (tmp_path / "sales.csv").write_text("Name,Revenue\nAlice,4200\nBob,9100\n")
        provider = DocumentKnowledgeProvider(tmp_path)
        result = provider.lookup("Alice")
        assert result.excerpts
        assert "Alice" in result.excerpts[0].content
        assert "4200" in result.excerpts[0].content

    def test_all_rows_reach_evidence_for_an_aggregate_question(self, tmp_path):
        """The real, stated boundary this checkpoint documents: an
        aggregate question ("highest value") shares no words with the
        data, so every row - not just keyword matches - is offered,
        bounded by SPREADSHEET_MAX_EXCERPTS rather than excluded."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Name", "Revenue"])
        for i, name in enumerate(["Alice", "Bob", "Carol"]):
            ws.append([name, 1000 * (i + 1)])
        wb.save(tmp_path / "sales.xlsx")

        provider = DocumentKnowledgeProvider(tmp_path)
        result = provider.lookup("what is the highest value")
        contents = " ".join(e.content for e in result.excerpts)
        assert "Alice" in contents and "Bob" in contents and "Carol" in contents

    def test_an_unrelated_document_is_unaffected_by_the_wider_spreadsheet_cap(self, tmp_path):
        (tmp_path / "notes.md").write_text("Unrelated project notes.\n\nMore unrelated notes.")
        provider = DocumentKnowledgeProvider(tmp_path)
        result = provider.lookup("unrelated project")
        assert len(result.excerpts) <= 3  # ordinary MAX_EXCERPTS, not widened

    def test_the_full_production_path_answers_a_real_spreadsheet_question(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Name", "Revenue"])
        ws.append(["Alice", 4200])
        ws.append(["Bob", 9100])
        ws.append(["Carol", 3050])
        wb.save(tmp_path / "sales.xlsx")

        settings = Settings.from_env(
            env={}, memory="ephemeral", knowledge_provider="documents", knowledge_dir=str(tmp_path)
        )
        runtime, manager, channel, conversation = build(settings)
        _with_brain(runtime, manager, "The highest revenue is 9100, from Bob.")

        turn = exchange(channel, conversation, "read this spreadsheet and tell me what the highest value is")
        assert turn.succeeded
        assert "9100" in turn.outcome

    def test_a_malformed_spreadsheet_does_not_take_down_the_whole_lookup(self, tmp_path):
        (tmp_path / "broken.xlsx").write_bytes(b"not actually a spreadsheet")
        (tmp_path / "notes.md").write_text("Project Beacon planning notes are here.")
        provider = DocumentKnowledgeProvider(tmp_path)
        result = provider.lookup("Project Beacon planning")
        assert result.excerpts
        assert result.excerpts[0].document == "notes.md"


# ---------------------------------------------------------------------------
# 2. Presentation support - real slide text
# ---------------------------------------------------------------------------


class TestPresentationSupport:
    def test_slide_text_is_extracted_with_slide_locators(self, tmp_path):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Q3 Results"
        slide.placeholders[1].text = "Revenue grew 20 percent this quarter."
        prs.save(tmp_path / "deck.pptx")

        provider = DocumentKnowledgeProvider(tmp_path)
        result = provider.lookup("Q3 results revenue")
        assert result.excerpts
        assert "Q3 Results" in result.excerpts[0].content
        assert result.excerpts[0].locator == "slide 1"

    def test_which_slide_discusses_x_finds_the_right_slide(self, tmp_path):
        prs = Presentation()
        for title, body in [("Intro", "Welcome to the deck."), ("Budget", "The budget is finalised.")]:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = body
        prs.save(tmp_path / "deck.pptx")

        provider = DocumentKnowledgeProvider(tmp_path)
        result = provider.lookup("budget finalised")
        assert result.excerpts
        assert "Budget" in result.excerpts[0].content
        assert result.excerpts[0].locator == "slide 2"

    def test_the_full_production_path_summarises_a_real_presentation(self, tmp_path):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Q3 Results"
        slide.placeholders[1].text = "Revenue grew 20 percent this quarter."
        prs.save(tmp_path / "deck.pptx")

        settings = Settings.from_env(
            env={}, memory="ephemeral", knowledge_provider="documents", knowledge_dir=str(tmp_path)
        )
        runtime, manager, channel, conversation = build(settings)
        _with_brain(runtime, manager, "The deck reports 20% revenue growth in Q3.")

        turn = exchange(channel, conversation, "read this presentation and summarize it")
        assert turn.succeeded
        assert "20%" in turn.outcome or "20 percent" in turn.outcome.lower() or "growth" in turn.outcome.lower()

    def test_a_malformed_presentation_does_not_take_down_the_whole_lookup(self, tmp_path):
        (tmp_path / "broken.pptx").write_bytes(b"not actually a presentation")
        (tmp_path / "notes.md").write_text("Project Beacon planning notes are here.")
        provider = DocumentKnowledgeProvider(tmp_path)
        result = provider.lookup("Project Beacon planning")
        assert result.excerpts
        assert result.excerpts[0].document == "notes.md"


# ---------------------------------------------------------------------------
# 3. Image analysis - real mechanism, honest boundary
# ---------------------------------------------------------------------------


class TestImageAnalysisMechanism:
    def test_anthropic_provider_constructs_the_real_vision_request(self):
        """Mechanism-verified: a real (injected) Anthropic client double,
        never a live network call - the identical dependency-injection
        seam this project already uses for every other Anthropic test."""

        captured = {}

        class FakeMessages:
            def create(self, **kwargs):
                captured.update(kwargs)

                class Block:
                    type = "text"
                    text = "I see a red square."

                class Usage:
                    input_tokens = 100
                    output_tokens = 10

                class Message:
                    content = [Block()]
                    model = "claude-opus-5"
                    stop_reason = "end_turn"
                    usage = Usage()

                return Message()

        class FakeClient:
            messages = FakeMessages()

        provider = AnthropicProvider(api_key="test-key", client=FakeClient())
        completion = provider.analyze_image(b"\x89PNG-fake-bytes", "what do you see?")

        assert completion.text == "I see a red square."
        content = captured["messages"][0]["content"]
        assert content[0]["type"] == "image"
        assert content[0]["source"]["media_type"] == "image/png"
        assert content[1]["type"] == "text"
        assert content[1]["text"] == "what do you see?"

    def test_an_empty_prompt_or_empty_image_is_refused_before_any_call(self):
        provider = AnthropicProvider(api_key="k", client=object())
        with pytest.raises(ValueError):
            provider.analyze_image(b"bytes", "")
        with pytest.raises(ValueError):
            provider.analyze_image(b"", "describe this")

    def test_openai_and_null_providers_have_no_analyze_image_method(self):
        """The honest boundary, checked structurally: only the one
        provider that genuinely implements real vision support has the
        method at all."""
        from src.providers.null_provider import NullProvider
        from src.providers.openai_provider import OpenAIProvider

        assert not hasattr(NullProvider(), "analyze_image")
        assert not hasattr(OpenAIProvider(api_key="k"), "analyze_image")


class TestImageAnalysisCapability:
    def test_routing_recognises_the_phases_own_literal_example(self):
        text = "Alfred, look at this image photo.png and tell me what you see"
        assert image_analysis_intent(Task(description=text))
        cap = select_capability(Task(description=text), None, participants=["Alfred"])
        assert cap.name == "image_analysis"

    def test_a_reordered_request_also_routes_correctly(self):
        """Closure campaign (2026-09-16, N-06): the identical shared
        mechanism WS-02/ORG-01/create_artifact/create_diagram already
        use - "I want a look at this picture, can you describe it?"
        states the identical ask IMAGE_ANALYSIS_MARKERS' own fixed
        phrases cover, in the equally natural opposite order."""
        for text in (
            "I want a look at this picture, can you describe it?",
            "I have a photo: can you describe it?",
        ):
            assert image_analysis_intent(Task(description=text)), text
            cap = select_capability(Task(description=text), None, participants=["Alfred"])
            assert cap.name == "image_analysis", text

    def test_an_unrelated_declarative_statement_is_not_mistaken_for_an_image_request(self):
        assert not image_analysis_intent(Task(description="I have a meeting at 3pm."))

    def test_a_creation_request_does_not_falsely_claim_analysis_on_the_shared_noun(self):
        """Real word/intent collision found live (2026-09-16): "image"/
        "picture" are legitimate nouns for both an analysis request and a
        generation request - the noun alone previously let image_analysis
        (alphabetically first) silently win every such collision, even
        ones clearly asking to *create* something. The confirmation
        clause's own verb now disambiguates."""
        assert not image_analysis_intent(
            Task(description="I would love an image of a cat, can you make one?")
        )
        assert not image_analysis_intent(
            Task(description="I want a picture of a sunset, can you generate it?")
        )
        assert image_analysis_intent(
            Task(description="I want a look at this picture, can you describe it?")
        )

    def test_with_the_real_default_null_brain_refuses_honestly(self, tmp_path):
        """The real default system's own NullBrain still provides an
        honest analyze_image (bound, never None, exactly like converse)
        that refuses plainly rather than fabricating a description."""
        image_path = tmp_path / "photo.png"
        image_path.write_bytes(_real_png())
        runtime, manager, channel, conversation = real_system()
        turn = exchange(channel, conversation, f"look at this image {image_path} and tell me what you see")
        assert not turn.succeeded
        assert "no reasoning capability is configured" in turn.outcome

    def test_with_analyze_image_genuinely_unset_refuses_honestly(self):
        """The other, structural half: a caller that never wired the
        field at all (an older construction, or a bare test double) -
        request.analyze_image is None, checked directly against the
        handler."""
        from src.manager.capabilities import CapabilityRequest, image_analysis

        request = CapabilityRequest(
            task=Task(description="look at this image photo.png and tell me what you see"),
            identity=real_system()[0].identity,
            roster=lambda: (),
            own_name="Alfred",
            own_capabilities=(),
            analyze_image=None,
        )
        result = image_analysis(request)
        assert not result.succeeded
        assert "no vision-capable reasoning" in result.error.lower()

    def test_with_a_non_vision_provider_refuses_honestly_never_fabricating(self, tmp_path):
        """A real LLMManagerBrain wrapping a provider with no
        analyze_image method (the null/default case) - Brain-level
        honesty, never a Tool-level one."""
        image_path = tmp_path / "photo.png"
        image_path.write_bytes(_real_png())
        runtime, manager, channel, conversation = real_system()
        from src.providers.null_provider import NullProvider

        _with_vision_brain(runtime, manager, NullProvider())
        turn = exchange(channel, conversation, f"look at this image {image_path} and tell me what you see")
        assert not turn.succeeded
        assert "cannot perform visual reasoning" in turn.outcome

    def test_with_a_real_vision_capable_double_produces_a_genuine_analysis(self, tmp_path):
        image_path = tmp_path / "photo.png"
        real_bytes = _real_png()
        image_path.write_bytes(real_bytes)

        runtime, manager, channel, conversation = real_system()
        vision = _FakeVisionProvider("I see a small red square.")
        _with_vision_brain(runtime, manager, vision)

        turn = exchange(channel, conversation, f"look at this image {image_path} and tell me what you see")
        assert turn.succeeded
        assert turn.outcome == "I see a small red square."
        assert vision.calls and vision.calls[0][0] == real_bytes

    def test_a_missing_image_file_fails_honestly(self):
        runtime, manager, channel, conversation = real_system()
        _with_vision_brain(runtime, manager, _FakeVisionProvider())
        turn = exchange(channel, conversation, "look at this image nonexistent.png and tell me what you see")
        assert not turn.succeeded

    def test_no_image_path_referenced_fails_honestly(self):
        runtime, manager, channel, conversation = real_system()
        _with_vision_brain(runtime, manager, _FakeVisionProvider())
        turn = exchange(channel, conversation, "look at this image and tell me what you see")
        assert not turn.succeeded
        assert "don't see an image" in turn.outcome.lower()

    def test_no_tool_ever_holds_reasoning_machinery(self):
        """The architectural boundary this checkpoint's own repair
        restores, checked directly: nothing under src/agents/tools/
        imports src.providers.base - image analysis reaches the Brain,
        never a Tool holding an LLMProvider."""
        import ast
        from pathlib import Path as _Path

        root = _Path(__file__).resolve().parent.parent / "src" / "agents" / "tools"
        for path in sorted(root.rglob("*.py")):
            tree = ast.parse(path.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                if isinstance(node, ast.ImportFrom) and node.module:
                    assert node.module != "src.providers.base", path.name


# ---------------------------------------------------------------------------
# 4. Control Center - the real capability matrix
# ---------------------------------------------------------------------------


class TestControlCenterCapabilityMatrix:
    def test_every_capability_reflects_real_runtime_state_not_code_existence(self):
        runtime, manager, channel, conversation = real_system()
        centre = ControlCenter(runtime, manager)
        report = render_capabilities(centre)
        # Half Brain, null providers by default - every Brain- or
        # vendor-provider-dependent capability must honestly say
        # UNAVAILABLE, never READY merely because the class exists.
        assert "General Knowledge" in report and "UNAVAILABLE" in report
        for row in ("General Knowledge", "Live Information", "Research", "Files:",
                    "Image Analysis", "Document Generation", "AI Image Generation",
                    "Narration", "Video Composition", "AI Video Generation",
                    "Video Understanding (transcription)",
                    "Voice (speech I/O)", "Connection"):
            line = next(l for l in report.splitlines() if row in l)
            assert "UNAVAILABLE" in line, line
        # Avatar/Environment/Memory are genuinely, honestly READY even
        # with no Brain and no vendor configured at all - each is a
        # real, deterministic, no-credential, no-network capability
        # (the "default" providers), never a false positive. Multimodal
        # workspace/artifact checkpoint: Artifact Workspace is genuinely
        # READY too (Half-Brain safe given already-structured content -
        # readiness reflects the tool grant alone, not brain.is_full_brain),
        # and so is Document Rendering (PDF) - "simple_pdf" is real,
        # local, zero-credential, and the default, the same reasoning
        # Avatar/Environment's own "default" providers already establish.
        for row in ("Avatar", "Environment", "Memory", "Artifact Workspace (documents, images/diagrams, audio, video)", "Document Rendering (PDF)"):
            line = next(l for l in report.splitlines() if row in l)
            assert "READY" in line, line
        assert report.count("READY") == 5

    def test_image_analysis_becomes_ready_only_with_a_real_vision_capable_provider(self):
        runtime, manager, channel, conversation = real_system(
            provider="anthropic", api_key="sk-fake-for-mechanism-check"
        )
        centre = ControlCenter(runtime, manager)
        report = render_capabilities(centre)
        lines = [l for l in report.splitlines() if "Image Analysis" in l]
        assert lines and "READY" in lines[0]
        assert "supports vision" in lines[0]

    def test_the_matrix_is_reachable_through_the_real_command_dispatcher(self):
        runtime, manager, channel, conversation = real_system()
        centre = ControlCenter(runtime, manager)
        output = run_command(centre, "capabilities", 100, channel, conversation)
        assert "CAPABILITIES" in output


# ---------------------------------------------------------------------------
# 5. Control Center - real provider management, reusing the existing
#    Ecosystem registry and file-backed CredentialStore.
# ---------------------------------------------------------------------------


class TestControlCenterProviderManagement:
    def test_provider_with_no_args_lists_every_real_category(self):
        runtime, manager, channel, conversation = real_system()
        centre = ControlCenter(runtime, manager)
        output = run_provider_command(centre, "")
        assert "llm_provider" in output
        assert "memory" in output
        assert "knowledge_provider" in output

    def test_unknown_category_is_refused_with_a_helpful_list(self):
        runtime, manager, channel, conversation = real_system()
        centre = ControlCenter(runtime, manager)
        output = run_provider_command(centre, "not-a-real-category")
        assert "no such category" in output

    def test_add_select_revoke_round_trip_through_the_real_control_center(self, tmp_path, monkeypatch):
        # Universal provider/connection management repair: accounts are
        # now persisted keyed by *backend* ("anthropic"), never by
        # *category* ("llm_provider") - the key build_provider itself
        # reads (account.provider decides which builder to call). The
        # old category-keyed convention this test pinned would have
        # raised "Unknown provider 'llm_provider'" the moment it was
        # ever actually threaded into a real construction - exactly the
        # disconnect this repair closes.
        import src.services.credentials.account_store_file as account_store_file
        import scripts.control_center as control_center_module

        accounts_path = tmp_path / "accounts.json"
        monkeypatch.setattr(control_center_module, "DEFAULT_ACCOUNTS_PATH", accounts_path)

        runtime, manager, channel, conversation = real_system()
        centre = ControlCenter(runtime, manager)

        added = run_provider_command(centre, "llm_provider add anthropic work sk-test-123")
        assert "added" in added
        assert "sk-test-123" not in added  # never echoed back

        accounts = run_provider_command(centre, "llm_provider accounts")
        assert "work" in accounts
        assert "sk-test-123" not in accounts

        added2 = run_provider_command(centre, "llm_provider add anthropic personal sk-other-456")
        assert "added" in added2
        selected = run_provider_command(centre, "llm_provider select anthropic personal")
        assert "now active" in selected

        store = account_store_file.load_credential_store(accounts_path)
        assert store.active("anthropic").account_id == "personal"

        revoked = run_provider_command(centre, "llm_provider revoke anthropic personal")
        assert "revoked" in revoked
        store_after = account_store_file.load_credential_store(accounts_path)
        assert store_after.active("anthropic") is None
        assert store_after.get("anthropic", "work") is not None

    def test_add_refuses_a_backend_the_category_does_not_recognise(self, tmp_path, monkeypatch):
        import scripts.control_center as control_center_module

        monkeypatch.setattr(control_center_module, "DEFAULT_ACCOUNTS_PATH", tmp_path / "accounts.json")
        runtime, manager, channel, conversation = real_system()
        centre = ControlCenter(runtime, manager)

        result = run_provider_command(centre, "llm_provider add not-a-real-vendor work sk-test-123")
        assert "no such backend" in result
        assert "not-a-real-vendor" in result

    def test_validate_reports_mechanism_only_never_claiming_live_verification(self):
        runtime, manager, channel, conversation = real_system()
        centre = ControlCenter(runtime, manager)
        output = run_provider_command(centre, "llm_provider validate null")
        assert "OK" in output
        assert "no live external call was made" in output

    def test_validate_a_missing_credential_fails_honestly(self):
        runtime, manager, channel, conversation = real_system()
        centre = ControlCenter(runtime, manager)
        output = run_provider_command(centre, "llm_provider validate anthropic")
        assert "FAILED" in output


# ---------------------------------------------------------------------------
# 6. Full Brain / Half Brain boundaries, reconfirmed across every new path
# ---------------------------------------------------------------------------


class TestHalfBrainHonesty:
    def test_spreadsheet_reasoning_requires_full_brain(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Name", "Revenue"])
        ws.append(["Alice", 4200])
        wb.save(tmp_path / "sales.xlsx")

        settings = Settings.from_env(
            env={}, memory="ephemeral", knowledge_provider="documents", knowledge_dir=str(tmp_path)
        )
        runtime, manager, channel, conversation = build(settings)
        turn = exchange(channel, conversation, "read this spreadsheet and tell me what the highest value is")
        assert not turn.succeeded
        assert "no reasoning" in turn.outcome.lower() or "no language model" in turn.outcome.lower()

    def test_document_generation_requires_full_brain(self, tmp_path):
        runtime, manager, channel, conversation = real_system(documents_dir=str(tmp_path))
        exchange(channel, conversation, "some prior discussion")
        turn = exchange(channel, conversation, "create a report from these findings")
        assert not turn.succeeded


# ---------------------------------------------------------------------------
# 7. Extensibility (Section 14) - a genuinely unseen Phase 3 completion
#    category enters through the identical Ecosystem boundary.
# ---------------------------------------------------------------------------


class TestExtensibility:
    def test_a_never_before_seen_document_format_category_enters_through_the_identical_boundary(self):
        import subprocess

        result = subprocess.run(
            ["git", "grep", "-il", "ocr_extraction_provider", "--", "src", "scripts"],
            cwd=".", capture_output=True, text=True,
        )
        assert result.stdout.strip() == "", "the fixture category must be genuinely unseen"

        class FakeOcrEngine:
            def __init__(self, key, credential):
                if key == "cloud-ocr" and credential is None:
                    raise RuntimeError("no API key configured for cloud-ocr")

        registry = EcosystemRegistry()
        registry.register(
            GenericCategoryAdapter(
                "ocr_extraction_provider",
                list_backends=lambda: ["null", "local-tesseract", "cloud-ocr"],
                active_key=lambda: "null",
                build=lambda key, credential: FakeOcrEngine(key, credential),
                requires_credential=lambda key: key == "cloud-ocr",
                dependents=("scanned-PDF and image-document text extraction",),
            )
        )

        assert "ocr_extraction_provider" in registry.categories()
        adapter = registry.adapter_for("ocr_extraction_provider")
        assert not adapter.validate("cloud-ocr").ok
        assert adapter.validate("cloud-ocr", credential="fake-but-present").ok
        assert adapter.validate("local-tesseract").ok

        real_registry = build_ecosystem_registry(Settings.from_env(env={}))
        real_registry.register(adapter)
        assert real_registry.adapter_for("knowledge_provider") is not None
        assert real_registry.adapter_for("ocr_extraction_provider") is adapter
